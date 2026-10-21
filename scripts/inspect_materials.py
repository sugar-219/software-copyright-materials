"""Audit source fidelity, DOCX layout, PDF fonts and pages; render every page."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw
from pypdf import PdfReader
import pypdfium2 as pdfium

from materials import FILES, read_json, resolve_source, write_json
from business_review import business_pending


def font_info(value):
    obj = value.get_object()
    name = str(obj.get("/BaseFont", ""))
    if "/DescendantFonts" in obj:
        child = obj["/DescendantFonts"][0].get_object()
        descriptor = child.get("/FontDescriptor")
    else:
        descriptor = obj.get("/FontDescriptor")
    embedded = False
    if descriptor:
        descriptor = descriptor.get_object()
        embedded = any(key in descriptor for key in ["/FontFile", "/FontFile2", "/FontFile3"])
    return {"name": name, "embedded": embedded}


def inspect_docx(path, errors):
    doc = Document(path)
    for section in doc.sections:
        if abs(section.page_width.mm - 210) > 0.2 or abs(section.page_height.mm - 297) > 0.2:
            errors.append(path.name + "：页面不是 A4")
    roots = [doc.element.body]
    for section in doc.sections:
        roots.extend([section.header._element, section.footer._element])
    paragraphs = [Paragraph(p, doc) for root in roots for p in root.xpath(".//w:p")]
    allowed = {"宋体", "微软雅黑", "Consolas", "SimSun", "Microsoft YaHei"}
    for p in paragraphs:
        for r in p.runs:
            if not r.text.strip():
                continue
            props = r._element.rPr
            faces = props.find(qn("w:rFonts")) if props is not None else None
            if faces is None:
                errors.append(path.name + "：文本没有显式字体设置")
                continue
            for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
                if faces.get(qn("w:" + key)) not in allowed:
                    errors.append(path.name + "：出现缺失或非约定字体")
                    break
    return doc


def inspect_pdf(path, render_dir, errors, pending):
    reader = PdfReader(path)
    fonts = {}
    for number, page in enumerate(reader.pages, 1):
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        if abs(width - 595.276) > 1.5 or abs(height - 841.89) > 1.5:
            errors.append(f"{path.name} 第{number}页不是 A4")
        resources = page.get("/Resources", {}).get_object()
        for item in resources.get("/Font", {}).get_object().values() if "/Font" in resources else []:
            info = font_info(item)
            fonts[info["name"]] = info
        if "\ufffd" in (page.extract_text() or ""):
            pending.append(f"{path.name} 第{number}页提取文本出现替代字符，需核对显示和编码")
    allowed = ["simsun", "song", "宋体", "microsoftyahei", "msyh", "微软雅黑", "consolas"]
    for info in fonts.values():
        normalized = re.sub(r"[^a-z\u4e00-\u9fff]", "", info["name"].lower())
        if not any(name in normalized for name in allowed):
            errors.append(path.name + "：PDF出现非约定字体 " + info["name"])
        if not info["embedded"]:
            pending.append(path.name + "：PDF字体未确认内嵌 " + info["name"])
    render_dir.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(path))
    thumbnails = []
    try:
        for number in range(len(pdf)):
            page = pdf[number]
            textpage = page.get_textpage()
            width, height = page.get_size()
            for index in range(textpage.count_chars()):
                left, bottom, right, top = textpage.get_charbox(index)
                if right > left and top > bottom and (left < -1 or bottom < -1 or right > width + 1 or top > height + 1):
                    errors.append(f"{path.name} 第{number + 1}页存在页外文字")
                    break
            textpage.close()
            bitmap = page.render(scale=1.4)
            img = bitmap.to_pil().convert("RGB")
            img.save(render_dir / f"page-{number + 1:03}.png")
            img.thumbnail((280, 396))
            canvas = Image.new("RGB", (300, 424), "#dddddd")
            canvas.paste(img, ((300 - img.width) // 2, 5))
            ImageDraw.Draw(canvas).text((10, 405), str(number + 1), fill="black")
            thumbnails.append(canvas)
            bitmap.close()
            page.close()
    finally:
        pdf.close()
    for start in range(0, len(thumbnails), 12):
        batch = thumbnails[start:start + 12]
        sheet = Image.new("RGB", (900, 424 * ((len(batch) + 2) // 3)), "white")
        for i, img in enumerate(batch):
            sheet.paste(img, ((i % 3) * 300, (i // 3) * 424))
        sheet.save(render_dir / f"contact-{start // 12 + 1:02}.jpg", quality=88)
    return {"pages": len(reader.pages), "fonts": list(fonts.values()), "renders": str(render_dir)}


def normalized(text):
    return re.sub(r"\s", "", text)


def body_lines(page, software, number, errors):
    lines = [line for line in (page.extract_text() or "").splitlines() if line.strip()]
    expected = normalized(software["name"] + software["version"] + str(number))
    if not lines or normalized(lines[0]) != expected:
        errors.append(f"鉴别材料第{number}页页眉名称、版本或连续页码不符")
        return lines
    return lines[1:]


def source_page_errors(lines, rows, number):
    errors = []
    if normalized("".join(lines)) != normalized("".join(row["text"] for row in rows)):
        errors.append(f"源程序PDF第{number}页全文与所选源码不一致（含遗漏、重复或额外正文）")
    return errors


def review_item(category, message, kind=None, field=None):
    key = category + ":" + str(kind) + ":" + str(field) + ":" + message
    return {"id": hashlib.sha256(key.encode()).hexdigest()[:16], "category": category,
            "message": message, "document": kind, "field": field, "status": "open"}


def apply_review(items, review, artifacts):
    """Resolutions are evidence-backed and bound to the exact document bytes."""
    for item in items:
        decision = review.get("items", {}).get(item["id"])
        if not decision:
            continue
        status = decision.get("status")
        if status not in ["resolved", "awaiting_applicant"] or not decision.get("evidence"):
            raise ValueError("每项人工处理须给出 resolved/awaiting_applicant 与 evidence")
        if status == "awaiting_applicant" and item["category"] not in ["privacy", "fact"]:
            raise ValueError("规则或视觉问题不能通过标为本人填写来跳过")
        if item["document"]:
            if item["document"] not in artifacts or not decision.get("artifact_hashes") or decision["artifact_hashes"] != artifacts[item["document"]]:
                raise ValueError("人工复核的文件哈希与当前材料不一致")
        item.update(status=status, evidence=decision["evidence"])
    return items


def inspect(work, review_path=None):
    work = Path(work).resolve(strict=True)
    data = read_json(work / "build-input.json")
    export = read_json(work / "word-export-report.json")
    build = read_json(work / "build-report.json")
    mapping = read_json(work / "source-map.json")
    output, project = Path(data["output_dir"]), Path(data["project_root"])
    errors, pending, pdfs, artifacts, page_details = [], [], {}, {}, {}
    reuse = read_json(work / "reuse.json") if (work / "reuse.json").exists() else {}
    previous = read_json(Path(reuse["previous_work"]) / "inspection.json") if reuse else {}
    business_path = work / "business-review.json"
    business_errors, business_warnings = business_pending(project, data, read_json(business_path) if business_path.exists() else None)
    errors.extend(business_errors)
    pending.extend(review_item("fact", message) for message in business_warnings)
    for item in mapping["files"]:
        path = resolve_source(project, item["path"])
        if hashlib.sha256(path.read_bytes()).hexdigest() != item["sha256"]:
            errors.append("项目源码在生成后发生变化：" + item["path"])
    expected_rows = [row["text"] for page in mapping["submitted_rows"] for row in page]
    source = inspect_docx(output / FILES["source"], errors)
    if [p.text for p in source.paragraphs] != expected_rows:
        errors.append("源程序Word正文与源码映射不一致")
    expected_pages = (list(range(1, mapping["full_pages"] + 1)) if mapping["full_pages"] <= 60
                      else list(range(1, 31)) + list(range(mapping["full_pages"] - 29, mapping["full_pages"] + 1)))
    if mapping["selected_pages"] != expected_pages:
        errors.append("源程序取页范围不符合一般交存的连续性")
    for item in export["documents"]:
        kind = item["kind"]
        if kind in artifacts:
            errors.append("导出报告存在重复材料：" + kind)
            continue
        artifacts[kind] = {form: hashlib.sha256(Path(item[form]).read_bytes()).hexdigest() for form in ["docx", "pdf"]}
        if item["kind"] != "source":
            inspect_docx(Path(item["docx"]), errors)
        if kind in reuse.get("kinds", []) and artifacts[kind] == reuse["artifact_hashes"][kind]:
            report = dict(previous["pdfs"][kind], reused=True)
        else:
            notices = []
            report = inspect_pdf(Path(item["pdf"]), work / "renders" / kind, errors, notices)
            pending.extend(review_item("visual", message, kind) for message in notices)
        pdfs[item["kind"]] = report
        if report["pages"] != item["pages"]:
            errors.append(item["kind"] + "：Word与PDF页数不同")
        if item["kind"] == "source" and report["pages"] != build["expected_source_pages"]:
            errors.append("源程序实际分页与取页计划不一致，需调整版式重新生成")
        if item["kind"] == "source" and report["pages"] == build["expected_source_pages"]:
            source_pdf = PdfReader(item["pdf"])
            for number, (page, rows) in enumerate(zip(source_pdf.pages, mapping["submitted_rows"]), 1):
                lines = body_lines(page, data["software"], number, errors)
                errors.extend(source_page_errors(lines, rows, number))
                nonblank = len([line for line in lines if line.strip()])
                original_lines = len({(row["path"], row["line"]) for row in rows if row["text"].strip()})
                page_details.setdefault("source", []).append({"page": number, "nonblank_display_lines": nonblank,
                    "original_nonblank_lines": original_lines, "mapped_display_rows": len(rows)})
                if nonblank < 50:
                    if number < len(source_pdf.pages):
                        errors.append(f"源程序PDF第{number}页非空显示行不足50，空行不能计数，需重新排版")
                    else:
                        pending.append(review_item("rule", "源程序末页不足50条非空显示行，需核对末页适用规则", "source"))
        if item["kind"] == "manual" and report["pages"] > 60:
            errors.append("说明书超过60页，当前脚本未对说明书自动取页；需按已确认规则处理并复核Word/PDF对应关系")
        if kind == "manual":
            reader = PdfReader(item["pdf"])
            for number, page in enumerate(reader.pages, 1):
                lines = body_lines(page, data["software"], number, errors)
                images = len(page.images)
                page_details.setdefault("manual", []).append({"page": number, "nonblank_display_lines": len(lines), "images": images})
                if len(lines) < 30:
                    if images or number == len(reader.pages):
                        pending.append(review_item("rule", f"说明书第{number}页正文不足30行；需核对图文页或末页适用规则", "manual"))
                    else:
                        errors.append(f"说明书第{number}页无图且正文不足30行，需核实内容与版式")
        pending.append(review_item("visual", "逐页查看原始PNG，核对字体、边界、截图与分页；记录所查看的全部页码", kind))
    expected = {FILES["source"], FILES["manual"], FILES["application"], "源程序鉴别材料.pdf", "软件说明书.pdf"}
    if data["software"]["development_mode"] == "合作开发":
        expected.add(FILES["agreement"])
    actual = {p.name for p in output.iterdir() if p.is_file() and not p.name.startswith("~$")}
    if actual != expected:
        errors.append("交付目录的文件集合不符合5/6文件约定")
    kinds = {"source", "manual", "application"} | ({"agreement"} if "合作开发协议模板.docx" in expected else set())
    if set(artifacts) != kinds:
        errors.append("导出报告没有完整覆盖本次材料")
    for issue in build["application_issues"]:
        pending.append(review_item(issue.get("category", "fact"), issue["issue"], "application", field=issue["field"]))
    for issue in build.get("agreement_issues", []):
        pending.append(review_item(issue.get("category", "fact"), issue["issue"], "agreement", field=issue["field"]))
    rules = data.get("rules", {})
    if not rules.get("checked_at") or not rules.get("sources"):
        pending.append(review_item("rule", "未记录本次官网规则核对日期或来源"))
    for source in rules.get("sources", []):
        domain = urlparse(source.get("url", "")).hostname or ""
        if not any(domain == base or domain.endswith("." + base) for base in ["ncac.gov.cn", "ccopyright.com.cn"]):
            pending.append(review_item("rule", "规则来源并非国家版权局或版权中心，需核对机构、适用范围和原文：" + source.get("url", "")))
    pending.extend(review_item("rule", message) for message in rules.get("pending", []))
    pending.append(review_item("fact", "核对真实功能、截图来源与用户确认的权属、日期；自动检查不代表登记通过"))
    pending = list({p["id"]: p for p in pending}.values())
    # Only matching-byte reused documents may carry forward prior human decisions.
    inherited = {p["id"]: p for p in reuse.get("review_items", []) if isinstance(p, dict)}
    for item in pending:
        previous_item = inherited.get(item["id"], {})
        old_status = previous_item.get("status")
        reusable_status = old_status == "resolved" or (old_status == "awaiting_applicant" and item["category"] in ["privacy", "fact"])
        if item["category"] != "rule" and item["document"] in reuse.get("kinds", []) and artifacts[item["document"]] == reuse["artifact_hashes"][item["document"]] and reusable_status:
            item.update(status=old_status, evidence=previous_item.get("evidence", "此前对应文件已核对"))
    if review_path:
        pending = apply_review(pending, read_json(review_path), artifacts)
    unresolved = [p for p in pending if p["status"] != "resolved"]
    report = {"technical_checks_passed": not errors, "errors": list(dict.fromkeys(errors)),
              "pending_review": pending, "pdfs": pdfs, "page_details": page_details, "artifact_hashes": artifacts,
              "unresolved_counts": {category: sum(p["category"] == category for p in unresolved) for category in ["privacy", "fact", "rule", "visual"]},
              "delivery_files": sorted(actual), "status": "technical_errors" if errors else ("draft_pending_review" if unresolved else "review_complete")}
    write_json(work / "inspection.json", report)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--work", required=True)
    parser.add_argument("--review", help="人工逐项处理记录JSON，须有证据及所核对材料哈希")
    args = parser.parse_args()
    report = inspect(args.work, args.review)
    print(json.dumps({"technical_checks_passed": report["technical_checks_passed"],
                      "errors": report["errors"], "pending_count": sum(p["status"] != "resolved" for p in report["pending_review"]),
                      "unresolved_counts": report["unresolved_counts"]}, ensure_ascii=False))
    return 0 if report["technical_checks_passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
