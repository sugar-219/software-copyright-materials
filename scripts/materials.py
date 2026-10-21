"""Create editable filing drafts from explicit project evidence on Windows."""
import argparse
import copy
import datetime
import hashlib
import json
import os
import re
import shutil
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageFont

from confirm_input import verify_confirmation
from business_review import require_business_review

ROOT = Path(__file__).resolve().parents[1]
FILES = {
    "manual": "软件说明书.docx",
    "source": "源程序鉴别材料.docx",
    "application": "申请信息填写稿.docx",
    "agreement": "合作开发协议模板.docx",
}
ROWS_PER_PAGE = 55
CODE_SIZE = 10
CODE_LEADING = 12


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def font(run_or_style, latin="宋体", east="宋体", size=11):
    run_or_style.font.name = latin
    run_or_style.font.size = Pt(size)
    run_or_style.font.color.rgb = RGBColor(0, 0, 0)
    props = run_or_style._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for key in list(fonts.attrib):
        del fonts.attrib[key]
    for key, value in {"ascii": latin, "hAnsi": latin, "cs": latin, "eastAsia": east}.items():
        fonts.set(qn("w:" + key), value)


def run(paragraph, text, latin="宋体", east="宋体", size=11):
    value = paragraph.add_run(str(text))
    font(value, latin, east, size)
    return value


def field(paragraph, instruction):
    value = run(paragraph, "", size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = " " + instruction + " "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    value._r.extend([begin, code, end])


def base_document(kind):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(18 if kind == "source" else 22)
    sec.left_margin = sec.right_margin = Mm(18 if kind == "source" else 23)
    sec.header_distance = sec.footer_distance = Mm(8)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Caption", "Header", "Footer"]:
        style = doc.styles[name]
        heading = name in ["Title", "Heading 1", "Heading 2"]
        font(style, "微软雅黑" if heading else "宋体", "微软雅黑" if heading else "宋体",
             {"Title": 18, "Heading 1": 14, "Heading 2": 12, "Caption": 9,
              "Header": 9, "Footer": 9}.get(name, 11))
        style.paragraph_format.space_before = Pt(8 if heading else 0)
        style.paragraph_format.space_after = Pt(5 if heading else 3)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.widow_control = True
        if heading:
            style.paragraph_format.keep_with_next = True
    style = doc.styles.add_style("SourceCode", WD_STYLE_TYPE.PARAGRAPH)
    font(style, "Consolas", "宋体", CODE_SIZE)
    fmt = style.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing_rule, fmt.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(CODE_LEADING)
    fmt.keep_with_next = fmt.keep_together = fmt.widow_control = False
    # Some bundled default templates carry a blue rule under Title/Subtitle.
    for border in list(doc.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    doc.styles["Header"].paragraph_format.tab_stops.clear_all()
    header = sec.header.paragraphs[0]
    header.paragraph_format.tab_stops.add_tab_stop(
        sec.page_width - sec.left_margin - sec.right_margin, WD_TAB_ALIGNMENT.RIGHT)
    run(header, "{{软件名称}}  {{版本号}}\t", size=9)
    field(header, "PAGE")
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    return doc


def add_title(doc, name, version, title):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, title, "微软雅黑", "微软雅黑", 18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, name + "  " + version)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading " + str(level))
    run(p, text, "微软雅黑", "微软雅黑", 14 if level == 1 else 12)
    return p


def paragraph(doc, text):
    p = doc.add_paragraph()
    run(p, text)
    return p


def party_label(index):
    return "甲乙丙丁戊己庚辛壬癸"[index] + "方" if index < 10 else f"合作方{index + 1}"


def cooperative_parties(cooperation):
    parties = cooperation.get("parties")
    if parties is not None:
        if not isinstance(parties, list) or len(parties) < 2 or any(not isinstance(p, dict) for p in parties):
            raise ValueError("合作方需为至少两项的 parties 对象列表")
        if "party_count" in cooperation and cooperation["party_count"] != len(parties):
            raise ValueError("合作人数与 parties 名单数量不一致")
        names = [p.get("name", "") for p in parties]
        known = [name for name in names if name and "【" not in name]
        if len(known) != len(set(known)):
            raise ValueError("合作方名单存在重复姓名，请核实")
        return parties
    count = cooperation.get("party_count")
    if type(count) is not int or count < 2:
        raise ValueError("先确认合作方名单，或为通用草稿明确指定 party_count，不能默认为两人")
    return [{} for _ in range(count)]


def agreement_issues(cooperation):
    issues = []
    for i, party in enumerate(cooperative_parties(cooperation), 1):
        for key, label, category in [("name", "姓名或名称", "fact"), ("identity_number", "证件号码", "privacy")]:
            if not party.get(key) or "【" in party[key]:
                issues.append({"field": f"合作方{i}{label}", "issue": "待本人填写或确认", "category": category})
    for key in ["development_period", "division", "ownership", "shares", "income_and_costs", "rights_exercise", "third_party", "representative", "authority", "other_terms"]:
        if not cooperation.get(key) or "【" in cooperation[key]:
            issues.append({"field": "协议." + key, "issue": "待各方确认", "category": "fact"})
    issues.append({"field": "本人签字及实际签署日期", "issue": "需本人完成签署", "category": "privacy"})
    return issues


def normalize_data(data):
    """Share confirmed facts; never derive identity or ownership from project code."""
    data = copy.deepcopy(data)
    software, app = data["software"], data.setdefault("application", {})
    for key in ["name", "version", "development_mode", "short_name"]:
        if key in software and key in app and software[key] != app[key]:
            raise ValueError("统一底稿的软件字段冲突：" + key)
    short = software.get("short_name", app.get("short_name", ""))
    if not isinstance(short, str):
        raise ValueError("short_name 必须为字符串；无简称使用空字符串或省略")
    software["short_name"] = app["short_name"] = short
    if app.get("publication_status") == "未发表":
        current = app.get("first_publication", "")
        if current and "【" not in current and "不适用" not in current:
            raise ValueError("未发表状态与首次发表信息冲突")
        app["first_publication"] = "不适用（未发表）"
    elif app.get("publication_status") == "已发表" and "first_publication_date" in app:
        place = app.get("first_publication_place", "【本人填写】")
        app["first_publication"] = f"首次发表日期：{app['first_publication_date']}；首次发表地点：{place}"
    def date(value):
        match = re.fullmatch(r"(\d{4})(?:年|-)(\d{1,2})(?:月|-)(\d{1,2})日?", value or "")
        return datetime.date(*map(int, match.groups())) if match else None
    completed, published = date(app.get("completion_date")), date(app.get("first_publication_date"))
    if completed and published and published < completed:
        raise ValueError("首次发表日期早于开发完成日期，请核实")
    if software["development_mode"] == "合作开发":
        cooperation = data.setdefault("cooperation", {})
        parties = cooperative_parties(cooperation)
        cooperation["party_count"] = len(parties)
        known = [p["name"] for p in parties if p.get("name") and "【" not in p["name"]]
        representative = cooperation.get("representative", "")
        if representative and "【" not in representative and len(known) == len(parties) and representative not in known:
            raise ValueError("办理代表不在已确认的合作方名单中")
        if known and len(known) == len(parties):
            owners = app.get("owners", "")
            positions = [owners.find(name) for name in known]
            if owners and (min(positions) < 0 or positions != sorted(positions)):
                raise ValueError("申请填写稿的权利人名单或顺序与合作方不一致")
            if not owners:
                app["owners"] = "\n\n".join(
                    f"著作权人{i + 1}：{p['name']}\n国家地区及身份类别：【本人填写】\n"
                    f"{('身份证号码' if p.get('identity_type') in ['居民身份证', '身份证'] else '证件类型及号码')}："
                    f"{p.get('identity_number', '【本人填写】')}\n联系信息：【按官方系统要求填写】"
                    for i, p in enumerate(parties))
            manual = data.get("manual")
            if manual is not None:
                if "authors" in manual and manual["authors"] != known:
                    raise ValueError("说明书署名与合作方名单不一致")
                manual["authors"] = known
    return data


def agreement_body(doc, name, version, party_count=2, cooperation=None, application=None):
    cooperation = cooperation or {"party_count": party_count}
    application = application or {}
    parties = cooperative_parties(cooperation)
    optional = set(cooperation.get("include_fields", []))
    allowed = {"identity_type", "address", "contact", "signing_place", "publication_place"}
    if optional - allowed:
        raise ValueError("未知协议可选字段：" + ",".join(sorted(optional - allowed)))
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Mm(20)
        sec.left_margin = sec.right_margin = Mm(22)
        for p in sec.header.paragraphs:
            p.clear()
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after, normal.line_spacing = Pt(3), Pt(15)
    normal.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(title, "软件合作开发协议", "微软雅黑", "微软雅黑", 18).bold = True
    for i, party in enumerate(parties):
        p = paragraph(doc, party_label(i) + "：" + party.get("name", "【本人填写】"))
        identity_block = [p]
        p.paragraph_format.tab_stops.add_tab_stop(Mm(46))
        label = "身份证号码" if party.get("identity_type") in ["身份证", "居民身份证"] else "证件号码"
        run(p, "\t" + label + "：" + party.get("identity_number", "【本人填写】"))
        p.paragraph_format.keep_together = True
        for key, label in [("identity_type", "证件类型"), ("address", "地址"), ("contact", "联系方式")]:
            if key in optional:
                identity_block.append(paragraph(doc, label + "：" + party.get(key, "【本人填写】")))
        for p in identity_block[:-1]:
            p.paragraph_format.keep_with_next = True
    paragraph(doc, "各方就以下软件的合作开发事实、著作权归属及登记办理事宜约定如下。")
    def section(title, texts):
        p = doc.add_paragraph(style="Heading 1")
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(3)
        run(p, title, "微软雅黑", "微软雅黑", 12).bold = True
        for text in texts:
            p = paragraph(doc, text)
            p.paragraph_format.first_line_indent = Pt(22)
    basics = f"软件名称：{name}；版本号：{version}。开发期间：{cooperation.get('development_period', '【各方确认后填写】')}。"
    if application.get("completion_date"):
        basics += "开发完成日期：" + application["completion_date"] + "。"
    if application.get("publication_status"):
        basics += "发表状态：" + application["publication_status"] + "。"
    if application.get("publication_status") == "已发表" and application.get("first_publication_date"):
        basics += "首次发表日期：" + application["first_publication_date"] + "。"
    if "publication_place" in optional:
        basics += "首次发表地点：" + application.get("first_publication_place", "【本人填写】") + "。"
    section("第一条 软件与开发期间", [basics])
    section("第二条 合作方式", [cooperation.get("division", "实际合作方式与分工：【各方确认后填写】")])
    section("第三条 著作权归属与权益", [
        "著作权归属：" + cooperation.get("ownership", "【各方确认后填写】"),
        "权益份额：" + cooperation.get("shares", "【各方协商后填写】") + "；收益及费用：" + cooperation.get("income_and_costs", "【各方协商后填写】"),
        "权利行使：" + cooperation.get("rights_exercise", "【各方协商后填写】"),
        "第三方代码与资源安排：" + cooperation.get("third_party", "【按实际使用情况确认】")])
    section("第四条 登记办理代表", [
        "办理代表：" + cooperation.get("representative", "【各方确认后填写】") + "；办理事项：" + cooperation.get("authority", "【各方确认后填写】"),
        "申请材料应与各方确认的权属及署名一致；应由本人签署的文件由本人签署，办理代表身份不构成著作权转让或代签授权。"])
    section("第五条 生效与其他约定", [
        cooperation.get("other_terms", "争议解决及其他约定：【各方协商后填写】"),
        f"本协议自各方全部签署之日起生效，一式{len(parties)}份，各方各执一份。签署日期按实际填写，不以开发完成日期替代。"])
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    signatures = []
    for start in range(0, len(parties), 2):
        group = parties[start:start + 2]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(86))
        for offset, party in enumerate(group):
            prefix = "\t" if offset else ""
            run(p, prefix + party_label(start + offset) + "签字或盖章：____________")
        p.paragraph_format.keep_with_next = True
        signatures.append(p)
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Mm(86))
        run(p, "\t".join("签署日期：____年__月__日" for _ in group))
        signatures.append(p)
    if len(parties) <= 4:
        for p in signatures[:-1]:
            p.paragraph_format.keep_with_next = True
    if "signing_place" in optional:
        paragraph(doc, "签署地点：【实际签署时填写】")


def initialize_templates():
    target = ROOT / "assets" / "templates"
    target.mkdir(parents=True, exist_ok=True)
    for kind, filename in FILES.items():
        doc = base_document(kind)
        if kind == "agreement":
            agreement_body(doc, "{{软件名称}}", "{{版本号}}")
        elif kind != "source":
            add_title(doc, "{{软件名称}}", "{{版本号}}", filename[:-5])
            paragraph(doc, "{{由项目真实内容填充正文}}")
        else:
            p = doc.add_paragraph(style="SourceCode")
            run(p, "{{从项目提取源程序}}", "Consolas", "宋体", CODE_SIZE)
        doc.save(target / filename)


def load_template(kind, software):
    doc = Document(ROOT / "assets" / "templates" / FILES[kind])
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for section in doc.sections:
        for p in section.header.paragraphs:
            for r in p.runs:
                if "{{软件名称}}" in r.text:
                    r.text = r.text.replace("{{软件名称}}", software["name"]).replace("{{版本号}}", software["version"])
                    font(r, size=9)
    return doc


class CodeWrapper:
    def __init__(self, font_dir=None):
        directory = Path(font_dir or Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts")
        self.latin = ImageFont.truetype(str(directory / "consola.ttf"), CODE_SIZE * 4)
        self.east = ImageFont.truetype(str(directory / "simsun.ttc"), CODE_SIZE * 4)
        # Leave 4 mm of safety for differences between Pillow and Word shaping.
        self.width = (210 - 36 - 4) / 25.4 * 72 * 4
        self.widths = {}

    def wrap(self, text):
        expanded = text.expandtabs(4)
        if not expanded:
            return [""]
        chunks, current, width = [], "", 0
        for char in expanded:
            if ord(char) > 0xFFFF or (ord(char) < 32):
                raise ValueError("源程序包含控制字符或尚未支持字体的补充平面字符，请先确认其显示方案")
            if char not in self.widths:
                face = self.latin if ord(char) < 128 else self.east
                self.widths[char] = face.getlength(char)
            advance = self.widths[char]
            if current and width + advance > self.width and not unicodedata.combining(char):
                chunks.append(current)
                current, width = "", 0
            current += char
            width += advance
        chunks.append(current)
        assert "".join(chunks) == expanded
        return chunks


def resolve_source(project, relative):
    path = Path(relative)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError("源码必须使用项目内相对路径：" + relative)
    resolved = (project / path).resolve(strict=True)
    if not resolved.is_relative_to(project):
        raise ValueError("源码链接指向项目外：" + relative)
    if not resolved.is_file():
        raise ValueError("源码不是普通文件：" + relative)
    return resolved


def collect_source(project, specs, wrapper, omit_blank_lines=True):
    rows, files, seen = [], [], set()
    if not specs:
        raise ValueError("必须明确列出本次纳入的源文件及顺序")
    secret = re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----|\bAKIA[0-9A-Z]{16}\b|\bsk-(?:proj-)?[A-Za-z0-9_-]{24,}")
    for spec in specs:
        relative = spec["path"]
        path = resolve_source(project, relative)
        key = str(path).casefold()
        if key in seen:
            raise ValueError("源文件重复：" + relative)
        seen.add(key)
        raw = path.read_bytes()
        text = raw.decode(spec.get("encoding", "utf-8-sig"), errors="strict")
        if secret.search(text):
            raise ValueError("源文件疑似含真实凭据，先处理，不输出敏感值：" + relative)
        original = re.split(r"\r\n|\n|\r", text) if text else []
        if original and original[-1] == "" and text.endswith(("\n", "\r")):
            original.pop()
        files.append({"path": relative, "encoding": spec.get("encoding", "utf-8-sig"),
                      "sha256": hashlib.sha256(raw).hexdigest(), "physical_lines": len(original),
                      "nonblank_lines": sum(bool(line.strip()) for line in original),
                      "omitted_blank_lines": [i for i, line in enumerate(original, 1) if omit_blank_lines and not line.strip()]})
        for number, line in enumerate(original, 1):
            if omit_blank_lines and not line.strip():
                continue
            for segment, content in enumerate(wrapper.wrap(line), 1):
                rows.append({"path": relative, "line": number, "segment": segment, "text": content})
    if not rows:
        raise ValueError("纳入的源码为空")
    return files, rows


def select_pages(rows, per_page=ROWS_PER_PAGE):
    pages = [rows[i:i + per_page] for i in range(0, len(rows), per_page)]
    if len(pages) > 1 and len(pages[-1]) < 50 and len(rows) >= len(pages) * 50:
        base, extra = divmod(len(rows), len(pages))
        sizes = [base + (i < extra) for i in range(len(pages))]
        pages, cursor = [], 0
        for size in sizes:
            pages.append(rows[cursor:cursor + size])
            cursor += size
    indices = list(range(len(pages))) if len(pages) <= 60 else list(range(30)) + list(range(len(pages) - 30, len(pages)))
    return pages, indices


def make_source(software, rows):
    doc = load_template("source", software)
    pages, indices = select_pages(rows)
    for position, page_index in enumerate(indices):
        for row_index, row in enumerate(pages[page_index]):
            p = doc.add_paragraph(style="SourceCode")
            if position and row_index == 0:
                p.paragraph_format.page_break_before = True
            run(p, row["text"], "Consolas", "宋体", CODE_SIZE)
    return doc, pages, indices


def make_manual(software, manual, data_dir):
    doc = load_template("manual", software)
    add_title(doc, software["name"], software["version"], "软件说明书")
    if manual.get("authors"):
        paragraph(doc, "著作权人：" + "、".join(manual["authors"]))
    evidence = []
    if not manual.get("sections"):
        raise ValueError("说明书需要由 agent 编写真实章节，不能使用占位说明代替")
    for section in manual["sections"]:
        heading(doc, section["heading"])
        if not section.get("blocks"):
            raise ValueError("说明书存在空章节：" + section["heading"])
        for block in section["blocks"]:
            if "text" in block:
                for text in block["text"].split("\n"):
                    if text.strip():
                        paragraph(doc, text)
            elif "image" in block:
                if block.get("origin") not in ["runtime_capture", "user_provided"]:
                    raise ValueError("截图必须标记为实际运行截取或用户提供")
                path = Path(block["image"])
                if not path.is_absolute():
                    path = data_dir / path
                with Image.open(path) as img:
                    width, height = img.size
                    size = min(164 / width, 130 / height)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = bool(block.get("caption"))
                p.add_run().add_picture(str(path), width=Mm(width * size), height=Mm(height * size))
                if block.get("caption"):
                    p = doc.add_paragraph(style="Caption")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run(p, block["caption"], size=9)
                evidence.append({"image": str(path.resolve()), "origin": block["origin"],
                                 "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
            else:
                raise ValueError("说明书块仅支持 text 或 image")
    return doc, evidence


def make_application(software, values, stats, limits=None):
    doc = load_template("application", software)
    add_title(doc, software["name"], software["version"], "申请信息填写稿")
    paragraph(doc, "本稿供复制填写与核对，不替代登记系统生成的正式申请表或签章页。")
    paragraph(doc, "【本人填写】为隐私留空项；【待本人确认】为尚需确认的事实。")
    values = dict(values)
    if values.get("publication_status") == "未发表":
        values.setdefault("first_publication", "不适用（未发表）")
    values.update(name=software["name"], version=software["version"],
                  development_mode=software["development_mode"], source_lines=str(stats["nonblank_lines"]))
    issues = []
    for group in read_json(ROOT / "assets" / "fields.json"):
        heading(doc, group["title"])
        for original_item in group["fields"]:
            item = dict(original_item)
            changes = (limits or {}).get(item["key"], {})
            if set(changes) - {"min", "max"} or any(type(v) is not int or v < 0 for v in changes.values()):
                raise ValueError("字段限制只接受非负整数 min/max：" + item["key"])
            item.update(changes)
            value = values.get(item["key"], item.get("default", "【待本人确认】"))
            if not isinstance(value, str):
                raise ValueError("申请信息必须为字符串：" + item["key"])
            if item.get("optional") and not value.strip():
                continue
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            run(p, item["label"] + "：").bold = True
            for block in value.split("\n\n"):
                paragraph(doc, block).paragraph_format.keep_together = True
            pending = (not value and not item.get("optional")) or "【" in value
            if pending:
                category = "privacy" if "本人填写" in value and "待本人确认" not in value else "fact"
                issues.append({"field": item["label"], "issue": "待填写或确认", "category": category})
            elif len(value) > item.get("max", float("inf")) or len(value) < item.get("min", 0):
                issues.append({"field": item["label"], "issue": "长度不符用户页面提示，需按当前页面确认", "length": len(value), "category": "rule"})
    return doc, issues


def build(project, data_path, output, work):
    project = Path(project).resolve(strict=True)
    data_path = Path(data_path).resolve(strict=True)
    output, work = Path(output).resolve(), Path(work).resolve()
    if output == work or output.is_relative_to(work) or work.is_relative_to(output):
        raise ValueError("交付目录和工作目录必须分开且不互相包含")
    if project == output or project.is_relative_to(output):
        raise ValueError("交付目录不能覆盖项目根或其上级")
    if output.exists() and any(output.iterdir()):
        raise ValueError("请提供空交付目录，避免覆盖已有材料")
    if work.exists() and any(work.iterdir()):
        raise ValueError("请提供空的单次工作目录，避免覆盖已有记录")
    if not project.is_dir():
        raise ValueError("项目路径不是目录")
    data = normalize_data(read_json(data_path))
    software = data["software"]
    if software.get("confirmed") is not True:
        raise ValueError("软件名称、申请版本和开发方式需由用户确认后置 confirmed=true")
    for key in ["name", "version", "development_mode"]:
        if not isinstance(software.get(key), str) or not software[key].strip() or "【" in software[key]:
            raise ValueError("缺少已确认的软件信息：" + key)
    if software["development_mode"] not in ["单独开发", "合作开发", "委托开发", "下达任务开发"]:
        raise ValueError("开发方式需明确")
    confirmation = verify_confirmation(project, data_path)
    business = require_business_review(project, data)
    files, rows = collect_source(project, data["source_files"], CodeWrapper(),
                                 data.get("source_layout", {}).get("omit_blank_lines", True))
    stats = {key: sum(item[key] for item in files) for key in ["physical_lines", "nonblank_lines"]}
    docs = {}
    docs["source"], pages, indices = make_source(software, rows)
    docs["manual"], images = make_manual(software, data["manual"], data_path.parent)
    docs["application"], issues = make_application(software, data.get("application", {}), stats,
                                                   data.get("field_limits"))
    if software["development_mode"] == "合作开发":
        docs["agreement"] = load_template("agreement", software)
        agreement_body(docs["agreement"], software["name"], software["version"],
                       cooperation=data["cooperation"], application=data["application"])
    output.mkdir(parents=True, exist_ok=True)
    work.mkdir(parents=True, exist_ok=True)
    for kind, doc in docs.items():
        doc.save(output / FILES[kind])
    mapping = {"files": files, "statistics": stats, "rows_per_page": ROWS_PER_PAGE,
               "full_pages": len(pages), "selected_pages": [i + 1 for i in indices],
               "submitted_rows": [pages[i] for i in indices]}
    write_json(work / "source-map.json", mapping)
    snapshot = dict(data)
    snapshot["project_root"] = str(project)
    snapshot["output_dir"] = str(output)
    snapshot["screenshot_evidence"] = images
    write_json(work / "build-input.json", snapshot)
    write_json(work / "input-confirmation.json", confirmation)
    write_json(work / "business-review.json", business)
    export = {"work_dir": str(work), "documents": []}
    for kind in docs:
        # Agreement and application are rendered to INTERNAL PDFs only for visual QA.
        pdf = output / (Path(FILES[kind]).stem + ".pdf") if kind in ["source", "manual"] else work / (kind + "-preview.pdf")
        export["documents"].append({"kind": kind, "docx": str(output / FILES[kind]), "pdf": str(pdf)})
    write_json(work / "export-job.json", export)
    write_json(work / "build-report.json", {"application_issues": issues,
               "agreement_issues": agreement_issues(data["cooperation"]) if "agreement" in docs else [],
               "rules": data.get("rules", {}), "expected_source_pages": len(indices),
               "expected_delivery_files": 6 if "agreement" in docs else 5,
               "status": "draft_pending_rule_content_and_visual_review"})
    return {"word_files": len(docs), "expected_source_pages": len(indices), "statistics": stats}


def revise(previous_work, data_path, output, work, documents):
    """Regenerate affected documents; reuse only artifacts with verified receipts."""
    previous_work = Path(previous_work).resolve(strict=True)
    data_path = Path(data_path).resolve(strict=True)
    output, work = Path(output).resolve(), Path(work).resolve()
    if output == work or output.is_relative_to(work) or work.is_relative_to(output):
        raise ValueError("交付目录与工作目录不能重合或互相包含")
    if output.is_relative_to(previous_work) or previous_work.is_relative_to(output) or work == previous_work or work.is_relative_to(previous_work) or previous_work.is_relative_to(work):
        raise ValueError("修订必须使用独立新目录，不覆盖此前记录")
    if any(p.exists() and any(p.iterdir()) for p in [output, work]):
        raise ValueError("修订交付目录和工作目录必须为空")
    old_data = read_json(previous_work / "build-input.json")
    data = normalize_data(read_json(data_path))
    project = Path(old_data["project_root"]).resolve(strict=True)
    confirmation = verify_confirmation(project, data_path)
    business = require_business_review(project, data)
    if project == output or project.is_relative_to(output):
        raise ValueError("修订交付不能覆盖项目根或其上级")
    previous = read_json(previous_work / "inspection.json")
    if not previous.get("technical_checks_passed") or not previous.get("artifact_hashes"):
        raise ValueError("复用前需使用新版检查器核验基线并记录 artifact_hashes；不可伪造既有检查结果")
    mapping = read_json(previous_work / "source-map.json")
    requested = set(documents)
    valid = {"source", "manual", "application"} | ({"agreement"} if data["software"]["development_mode"] == "合作开发" else set())
    if not requested or requested - valid:
        raise ValueError("documents 必须明确列出本次需要重制的材料类型")
    global_change = any(data["software"].get(k) != old_data["software"].get(k) for k in ["name", "version", "development_mode"])
    if global_change or data.get("source_files") != old_data.get("source_files") or data.get("source_layout") != old_data.get("source_layout"):
        raise ValueError("软件名称、版本、开发方式或源码范围变化，请执行完整 build 并重新确认申报范围")
    current_manual, previous_manual = copy.deepcopy(data.get("manual")), copy.deepcopy(old_data.get("manual"))
    for manual_data in [current_manual, previous_manual]:
        for section in (manual_data or {}).get("sections", []):
            section.pop("feature_ids", None)
    if current_manual != previous_manual:
        requested.add("manual")
    if data.get("application") != old_data.get("application") or data.get("field_limits") != old_data.get("field_limits") or data["software"].get("short_name") != old_data["software"].get("short_name"):
        requested.add("application")
        if "agreement" in valid:
            requested.add("agreement")
    if data.get("cooperation") != old_data.get("cooperation") and "agreement" in valid:
        requested.add("agreement")
        if data["cooperation"].get("parties") != old_data.get("cooperation", {}).get("parties"):
            requested.add("application")
    for entry in mapping["files"]:
        if hashlib.sha256(resolve_source(project, entry["path"]).read_bytes()).hexdigest() != entry["sha256"]:
            raise ValueError("源码已变化，停止局部复用；请重新核定版本并完整 build：" + entry["path"])
    if "source" in requested:
        raise ValueError("局部修订保留源程序；需要重新排版源程序时使用完整 build")
    docs = {}
    stats = mapping["statistics"]
    if "manual" in requested:
        docs["manual"], images = make_manual(data["software"], data["manual"], data_path.parent)
    else:
        images = old_data.get("screenshot_evidence", [])
    app_doc, issues = make_application(data["software"], data.get("application", {}), stats, data.get("field_limits"))
    if "application" in requested:
        docs["application"] = app_doc
    if "agreement" in requested:
        docs["agreement"] = load_template("agreement", data["software"])
        agreement_body(docs["agreement"], data["software"]["name"], data["software"]["version"],
                       cooperation=data["cooperation"], application=data["application"])
    previous_export = {i["kind"]: i for i in read_json(previous_work / "word-export-report.json")["documents"]}
    reuse = {}
    for kind in valid - requested:
        entry = previous_export[kind]
        receipt = previous["artifact_hashes"].get(kind, {})
        for form in ["docx", "pdf"]:
            path = Path(entry[form])
            if hashlib.sha256(path.read_bytes()).hexdigest() != receipt.get(form):
                raise ValueError("复用材料与基线核验记录不符：" + kind + "/" + form)
        reuse[kind] = entry
    output.mkdir(parents=True, exist_ok=True)
    work.mkdir(parents=True, exist_ok=True)
    job = {"work_dir": str(work), "documents": [], "reused_documents": []}
    for kind in sorted(valid):
        docx = output / FILES[kind]
        pdf = output / (Path(FILES[kind]).stem + ".pdf") if kind in ["source", "manual"] else work / (kind + "-preview.pdf")
        if kind in docs:
            docs[kind].save(docx)
            job["documents"].append({"kind": kind, "docx": str(docx), "pdf": str(pdf)})
        else:
            entry = dict(reuse[kind], docx=str(docx), pdf=str(pdf), reused=True)
            shutil.copy2(reuse[kind]["docx"], docx)
            shutil.copy2(reuse[kind]["pdf"], pdf)
            job["reused_documents"].append(entry)
    snapshot = dict(data, project_root=str(project), output_dir=str(output), screenshot_evidence=images)
    write_json(work / "build-input.json", snapshot)
    write_json(work / "input-confirmation.json", confirmation)
    write_json(work / "business-review.json", business)
    write_json(work / "source-map.json", mapping)
    write_json(work / "export-job.json", job)
    write_json(work / "reuse.json", {"previous_work": str(previous_work), "kinds": sorted(reuse),
                                   "artifact_hashes": {k: previous["artifact_hashes"][k] for k in reuse},
                                   "review_items": previous.get("pending_review", [])})
    write_json(work / "build-report.json", {"application_issues": issues, "rules": data.get("rules", {}),
               "agreement_issues": agreement_issues(data["cooperation"]) if "agreement" in valid else [],
               "expected_source_pages": len(mapping["selected_pages"]), "expected_delivery_files": 6 if "agreement" in valid else 5,
               "regenerated_documents": sorted(docs), "reused_documents": sorted(reuse), "status": "revision_pending_checks"})
    return {"regenerated_documents": sorted(docs), "reused_documents": sorted(reuse)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("init-templates")
    cmd = sub.add_parser("build")
    for argument in ["project", "data", "output", "work"]:
        cmd.add_argument("--" + argument, required=True)
    cmd = sub.add_parser("revise")
    for argument in ["previous-work", "data", "output", "work", "documents"]:
        cmd.add_argument("--" + argument, required=True)
    args = parser.parse_args()
    try:
        if args.command == "init-templates":
            initialize_templates()
            print("Templates created")
        elif args.command == "revise":
            print(json.dumps(revise(args.previous_work, args.data, args.output, args.work, args.documents.split(",")), ensure_ascii=False))
        else:
            print(json.dumps(build(args.project, args.data, args.output, args.work), ensure_ascii=False))
    except (ValueError, OSError, KeyError, UnicodeError) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
