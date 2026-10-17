"""Focused mechanical checks. Fixtures are not filing materials."""
import tempfile
import unittest
import argparse
import sys
from pathlib import Path

from docx import Document

from materials import (CodeWrapper, FILES, build, collect_source,
                       make_source, read_json, resolve_source, select_pages, write_json,
                       agreement_body, agreement_issues, base_document, font, make_application, normalize_data)
from inspect_materials import inspect_docx, source_page_errors, body_lines, review_item, apply_review


class MaterialsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wrapper = CodeWrapper()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.project = self.root / "project"
        self.project.mkdir()
        (self.project / "main.py").write_text("# 中文注释\n\ndef main():\n\tprint('验证缩进')\n", encoding="utf-8")
        self.data = {"software": {"name": "排版测试软件", "version": "2.3", "development_mode": "单独开发", "confirmed": True},
                     "source_files": [{"path": "main.py"}],
                     "manual": {"sections": [{"heading": "测试用途", "blocks": [{"text": "本材料仅验证排版与程序，不用于软著申请。"}]}]}}
        self.data_path = self.root / "input.json"
        write_json(self.data_path, self.data)

    def tearDown(self):
        self.tmp.cleanup()

    def test_source_path_escape_rejected(self):
        (self.root / "outside.py").write_text("print(1)")
        with self.assertRaises(ValueError):
            resolve_source(self.project.resolve(), "../outside.py")

    def test_duplicate_source_rejected(self):
        with self.assertRaises(ValueError):
            collect_source(self.project, [{"path": "main.py"}, {"path": "main.py"}], self.wrapper)

    def test_invalid_encoding_does_not_replace_characters(self):
        (self.project / "bad.py").write_bytes(b"\xff\xfe\xff")
        with self.assertRaises(UnicodeError):
            collect_source(self.project, [{"path": "bad.py"}], self.wrapper)

    def test_wrap_reconstructs_original_content(self):
        text = "\t" + "变量ABC_123 = '中文内容'; " * 80
        segments = self.wrapper.wrap(text)
        self.assertGreater(len(segments), 1)
        self.assertEqual("".join(segments), text.expandtabs(4))
        self.assertTrue(all(len(segment) > 0 for segment in segments))

    def test_page_boundaries_and_continuity(self):
        for count in [1, 59, 60, 61, 87]:
            rows = [{"text": str(i)} for i in range(count * 55)]
            pages, indices = select_pages(rows)
            expected = list(range(count)) if count <= 60 else list(range(30)) + list(range(count - 30, count))
            self.assertEqual(indices, expected)
            self.assertEqual(len(pages), count)
            self.assertEqual(len(set(indices)), len(indices))

    def test_original_and_displayed_line_counts_are_distinct(self):
        (self.project / "main.py").write_text("a = '" + "中" * 500 + "'\n\nprint(a)\n", encoding="utf-8")
        files, rows = collect_source(self.project, [{"path": "main.py"}], self.wrapper)
        self.assertEqual(files[0]["physical_lines"], 3)
        self.assertEqual(files[0]["nonblank_lines"], 2)
        self.assertGreater(len(rows), 3)

    def test_source_word_is_selected_original_rows(self):
        rows = [{"text": str(i)} for i in range(61 * 55)]
        doc, pages, indices = make_source(self.data["software"], rows)
        self.assertEqual([p.text for p in doc.paragraphs], [r["text"] for i in indices for r in pages[i]])
        self.assertEqual(sum(bool(p.paragraph_format.page_break_before) for p in doc.paragraphs), 59)

    def test_noncooperative_output_and_version(self):
        output, work = self.root / "delivery", self.root / "work"
        build(self.project, self.data_path, output, work)
        self.assertEqual({p.name for p in output.iterdir()}, {FILES[k] for k in ["source", "manual", "application"]})
        text = "\n".join(p.text for p in Document(output / FILES["application"]).paragraphs)
        self.assertIn("2.3", text)
        self.assertNotIn("1.0", text)
        job = read_json(work / "export-job.json")
        self.assertEqual(sum(Path(i["pdf"]).parent == output for i in job["documents"]), 2)

    def test_cooperative_word_only_and_no_assumed_shares(self):
        self.data["software"]["development_mode"] = "合作开发"
        self.data["cooperation"] = {"party_count": 3}
        write_json(self.data_path, self.data)
        output, work = self.root / "delivery", self.root / "work"
        build(self.project, self.data_path, output, work)
        doc = Document(output / FILES["agreement"])
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertEqual(text.count("签字或盖章："), 3)
        self.assertNotIn("平均", text)
        self.assertNotIn("50%", text)
        self.assertFalse((output / "合作开发协议模板.pdf").exists())
        agreement = next(i for i in read_json(work / "export-job.json")["documents"] if i["kind"] == "agreement")
        self.assertEqual(Path(agreement["pdf"]).parent, work)

    def test_existing_materials_not_overwritten(self):
        output = self.root / "delivery"
        output.mkdir()
        (output / "existing.txt").write_text("keep")
        with self.assertRaises(ValueError):
            build(self.project, self.data_path, output, self.root / "work")
        self.assertEqual((output / "existing.txt").read_text(), "keep")

    def test_unconfirmed_identity_information_cannot_be_guessed(self):
        self.data["software"]["confirmed"] = False
        write_json(self.data_path, self.data)
        with self.assertRaises(ValueError):
            build(self.project, self.data_path, self.root / "delivery", self.root / "work")

    def test_existing_work_records_not_overwritten(self):
        work = self.root / "work"
        work.mkdir()
        (work / "source-map.json").write_text("keep", encoding="utf-8")
        with self.assertRaises(ValueError):
            build(self.project, self.data_path, self.root / "delivery", work)
        self.assertEqual((work / "source-map.json").read_text(encoding="utf-8"), "keep")

    def test_missing_manual_image_fails_before_delivery(self):
        self.data["manual"]["sections"][0]["blocks"] = [{"image": "missing.png", "origin": "user_provided"}]
        write_json(self.data_path, self.data)
        output = self.root / "delivery"
        with self.assertRaises(FileNotFoundError):
            build(self.project, self.data_path, output, self.root / "work")
        self.assertFalse(output.exists())

    def cooperation(self, count=4):
        return {"parties": [{"name": f"测试人员{i}", "identity_type": "居民身份证"} for i in range(count)],
                "development_period": "2026年1月2日至2026年2月3日", "ownership": "共同享有全部著作权",
                "shares": "按各方另行确认的份额", "representative": "测试人员0"}

    def test_no_short_name_omits_entire_optional_field(self):
        doc, _ = make_application(self.data["software"], {}, {"nonblank_lines": 50})
        self.assertFalse(any("软件简称" in p.text for p in doc.paragraphs))
        doc, _ = make_application(self.data["software"], {"short_name": "测试简称"}, {"nonblank_lines": 50})
        self.assertIn("测试简称", "\n".join(p.text for p in doc.paragraphs))

    def test_confirmed_parties_and_facts_propagate(self):
        self.data["software"]["development_mode"] = "合作开发"
        self.data["cooperation"] = self.cooperation()
        self.data["application"] = {"publication_status": "已发表", "completion_date": "2026-02-03",
            "first_publication_date": "2026-02-05", "first_publication_place": "【本人填写】"}
        data = normalize_data(self.data)
        self.assertEqual(data["manual"]["authors"], [p["name"] for p in data["cooperation"]["parties"]])
        doc = base_document("agreement")
        agreement_body(doc, "排版测试软件", "2.3", cooperation=data["cooperation"], application=data["application"])
        text = "\n".join(p.text for p in doc.paragraphs)
        for value in ["测试人员0", "测试人员3", "2026-02-03", "2026-02-05", "共同享有全部著作权", "按各方另行确认的份额"]:
            self.assertIn(value, text)
        self.assertEqual(text.count("身份证号码："), 4)
        self.assertNotIn("签署地点", text)
        self.assertNotIn("联系方式", text)
        self.assertNotIn("首次发表地点", text)
        self.assertIn("首次发表地点", data["application"]["first_publication"])
        self.assertEqual(len([p for p in doc.paragraphs if "签字或盖章" in p.text]), 2)

    def test_identity_type_is_not_assumed(self):
        doc = base_document("agreement")
        cooperation = {"parties": [{"name": "测试甲"}, {"name": "测试乙", "identity_type": "护照"}]}
        agreement_body(doc, "排版测试软件", "2.3", cooperation=cooperation)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("身份证", text)
        self.assertEqual(text.count("证件号码："), 2)
        self.assertTrue(any(p.paragraph_format.keep_with_next for p in doc.paragraphs if "签字或盖章" in p.text))

    def test_agreement_optional_fields_are_explicit(self):
        doc = base_document("agreement")
        cooperation = self.cooperation(2)
        cooperation["include_fields"] = ["contact", "signing_place", "publication_place"]
        cooperation["parties"][0]["contact"] = "测试邮箱"
        agreement_body(doc, "排版测试软件", "2.3", cooperation=cooperation, application={"first_publication_place": "测试城市"})
        text = "\n".join(p.text for p in doc.paragraphs)
        for value in ["测试邮箱", "签署地点", "测试城市"]:
            self.assertIn(value, text)

    def test_cooperation_count_and_representative_conflicts_rejected(self):
        self.data["software"]["development_mode"] = "合作开发"
        self.data["cooperation"] = self.cooperation(2)
        self.data["cooperation"]["party_count"] = 4
        with self.assertRaises(ValueError):normalize_data(self.data)
        self.data["cooperation"].pop("party_count")
        self.data["cooperation"]["representative"] = "名单之外的人员"
        with self.assertRaises(ValueError):normalize_data(self.data)

    def test_unknown_party_count_does_not_default_to_two(self):
        self.data["software"]["development_mode"] = "合作开发"
        with self.assertRaises(ValueError):normalize_data(self.data)

    def test_conflicting_software_or_author_order_rejected(self):
        self.data["application"] = {"version": "9.9"}
        with self.assertRaises(ValueError):normalize_data(self.data)
        self.data["software"]["development_mode"] = "合作开发"
        self.data["cooperation"] = self.cooperation(2)
        self.data["application"] = {"owners": "测试人员1\n测试人员0"}
        with self.assertRaises(ValueError):normalize_data(self.data)

    def test_publication_state_and_dates_must_agree(self):
        self.data["application"] = {"publication_status": "未发表", "first_publication": "【本人填写】"}
        self.assertEqual(normalize_data(self.data)["application"]["first_publication"], "不适用（未发表）")
        self.data["application"] = {"completion_date": "2026-02-03", "first_publication_date": "2026-01-01"}
        with self.assertRaises(ValueError):normalize_data(self.data)

    def test_blank_lines_never_inflate_default_source_pagination(self):
        (self.project / "main.py").write_text("\n".join(f"v{i} = {i}\n\n\n" for i in range(100)), encoding="utf-8")
        files, rows = collect_source(self.project, [{"path": "main.py"}], self.wrapper)
        self.assertEqual(files[0]["nonblank_lines"], 100)
        self.assertTrue(files[0]["omitted_blank_lines"])
        self.assertEqual(len(rows), 100)
        pages, _ = select_pages(rows)
        self.assertEqual([len(p) for p in pages], [50, 50])
        _, preserved = collect_source(self.project, [{"path": "main.py"}], self.wrapper, omit_blank_lines=False)
        self.assertGreater(len(preserved), len(rows))

    def test_rebalanced_pages_preserve_all_rows_and_order(self):
        for count in [100, 150, 201, 301]:
            rows = [{"text": str(i)} for i in range(count)]
            pages, _ = select_pages(rows)
            self.assertEqual([r for p in pages for r in p], rows)
            if count >= len(pages) * 50:
                self.assertTrue(all(50 <= len(p) <= 55 for p in pages))

    def test_font_audit_includes_nested_table_cells(self):
        doc = base_document("manual")
        nested = doc.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
        font(nested.cell(0, 0).paragraphs[0].add_run("错误字体"), "Arial", "Arial")
        path = self.root / "table.docx";doc.save(path)
        errors = [];inspect_docx(path, errors)
        self.assertTrue(errors)

    def test_source_pdf_extra_or_missing_body_is_rejected(self):
        rows = [{"text": "print('first')"}, {"text": "print('last')"}]
        self.assertEqual(source_page_errors([r["text"] for r in rows], rows, 1), [])
        self.assertTrue(source_page_errors([r["text"] for r in rows] + ["unexpected()"], rows, 1))
        self.assertTrue(source_page_errors([rows[0]["text"]], rows, 1))

    def test_header_name_version_and_page_are_checked(self):
        class Page:
            def extract_text(self):return "错误名称 2.3 8\nprint(1)"
        errors = [];body_lines(Page(), self.data["software"], 1, errors)
        self.assertTrue(errors)

    def test_pending_fields_remain_distinct_and_need_evidence(self):
        a = review_item("privacy", "本人填写", "application", "甲方证件")
        b = review_item("privacy", "本人填写", "application", "乙方证件")
        self.assertNotEqual(a["id"], b["id"])
        review = {"items": {a["id"]: {"status": "awaiting_applicant", "evidence": "用户要求隐私留空", "artifact_hashes": {"docx": "hash"}}}}
        apply_review([a], review, {"application": {"docx": "hash"}})
        self.assertEqual(a["status"], "awaiting_applicant")
        with self.assertRaises(ValueError):apply_review([b], {"items": {b["id"]: {"status": "resolved", "evidence": "已看"}}}, {})
        rule = review_item("rule", "规则待核实")
        with self.assertRaises(ValueError):apply_review([rule], {"items": {rule["id"]: {"status": "awaiting_applicant", "evidence": "忽略"}}}, {})

    def test_missing_agreement_facts_are_classified_separately(self):
        issues = agreement_issues(self.cooperation(2))
        self.assertTrue(any(i["category"] == "privacy" for i in issues))
        self.assertTrue(any(i["category"] == "fact" for i in issues))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--temp-root")
    arguments, remaining = parser.parse_known_args()
    if arguments.temp_root:
        Path(arguments.temp_root).mkdir(parents=True, exist_ok=True)
        tempfile.tempdir = str(Path(arguments.temp_root).resolve())
    unittest.main(argv=[sys.argv[0], *remaining])
